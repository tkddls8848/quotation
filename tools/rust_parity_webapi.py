"""파이썬 API 층과 Rust 이식본을 같은 업로드로 대조한다.

브라우저에서 도는 것은 변환 코어만이 아니다. 파일 이름 다듬기, 크기·형식
거르기, 오류 문구, 응답 헤더, 구조화 로그까지 함께 돈다. 그 층을 Rust 로
옮겼으므로 (`rust/webapi`) 여기서 **계약이 같은지**를 확인한다.

    python tools/rust_parity_webapi.py

보는 것은 화면이 실제로 쓰는 것들이다.

    상태 코드 · 응답 헤더 · 오류 본문(코드·문구) · 구조화 로그 항목
    성공했을 때의 견적서 본문 (셀 단위)

두 자리만 값을 맞추지 않는다. `Content-Length` 는 ZIP 직렬화 차이로 파일
크기가 달라지므로 **제 본문 길이와 맞는지**를 보고, 구문 오류 문구의 뒷부분은
파서가 내는 설명이라 앞의 사용자 문구까지만 본다.

성공 본문은 `tools/rust_parity_workbook.py` 의 비교기를 그대로 쓴다. 두
라이브러리의 ZIP 직렬화 차이는 [결정 0005] 대로 중단 사유로 삼지 않는다.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
for path in (ROOT, ROOT / "web" / "src", ROOT / "tools"):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import openpyxl  # noqa: E402

import api  # noqa: E402
import template  # noqa: E402
import rust_parity_workbook as workbook  # noqa: E402

PROBE = ("cargo", "run", "--quiet", "--package", "quotation-parity",
         "--bin", "webapi-probe", "--")

FIXTURES = ROOT / "tests" / "fixtures" / "public"

#: 날짜를 고정해야 두 산출물이 같은 값을 갖는다.
TODAY = dt.date(2026, 8, 27)
DEPLOYMENT = "parity"
REQUEST_ID = "parity-request-id"

#: 통과해야 하는 업로드 — 이름, MIME, 내용.
def accepted() -> list[tuple[str, str, bytes]]:
    cases = [(path.name, "text/xml", path.read_bytes())
             for path in sorted(FIXTURES.glob("*.xml"))]
    # 브라우저가 MIME 을 못 붙이거나 다르게 붙이는 경우도 서버와 같아야 한다.
    sample = FIXTURES / "new_quote.xml"
    cases.append(("new_quote.xml", "", sample.read_bytes()))
    cases.append(("new_quote.xml", "application/octet-stream", sample.read_bytes()))
    # 경로가 붙은 이름, 한글 이름, 앞점 이름 — 다운로드 이름이 갈리는 자리다.
    cases.append((r"C:\Users\me\견 적 (2).xml", "text/xml", sample.read_bytes()))
    cases.append(("...견적서.xml", "text/xml", sample.read_bytes()))
    return cases


#: 거절해야 하는 업로드 — 이름, MIME, 내용.
REJECTED: list[tuple[str, str, bytes]] = [
    ("notes.txt", "text/plain", "XML 이 아닙니다".encode("utf-8")),
    ("quote.xml", "application/pdf", b"<CFXML/>"),
    ("empty.xml", "text/xml", b"   \n"),
    ("broken.xml", "text/xml", b"<CFXML><CFData><ProductLineItem>"),
    ("no-root.xml", "text/xml", b"<Other><CFData/></Other>"),
    ("no-items.xml", "text/xml", b"<CFXML><CFData/></CFXML>"),
    ("huge.xml", "text/xml", b"<CFXML><CFData>"
     + b"<ProductLineItem><TransactionType>NEW</TransactionType></ProductLineItem>" * 5001
     + b"</CFData></CFXML>"),
]


def python_response(name: str, content_type: str, content: bytes) -> dict:
    response = api.convert_response(
        [api.Upload(filename=name, content=content, content_type=content_type)],
        template_bytes=template.template_bytes,
        template_version=template.template_version,
        deployment_version=DEPLOYMENT,
        request_id=REQUEST_ID,
        today=TODAY)
    return {
        "status": response.status,
        "headers": dict(response.headers),
        # 시간은 기계마다 다르다. 대조 대상이 아니다.
        "log": {k: v for k, v in response.log.items() if k != "total_ms"},
        "body": response.body,
    }


def rust_response(name: str, content_type: str, content: bytes, body_path: Path) -> dict:
    result = subprocess.run(
        [*PROBE, name, content_type, DEPLOYMENT, REQUEST_ID, TODAY.isoformat(),
         str(body_path)],
        cwd=ROOT, input=content, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        check=False)
    if result.returncode != 0:
        sys.stderr.write(result.stderr.decode("utf-8", "replace"))
        raise SystemExit(f"Rust 탐침이 {result.returncode} 로 끝났습니다.")
    meta = json.loads(result.stdout.decode("utf-8"))
    return {
        "status": meta["status"],
        "headers": meta["headers"],
        "log": {k: v for k, v in meta["log"].items() if k != "total_ms"},
        "body": body_path.read_bytes(),
    }


def body_problems(label: str, want: dict, got: dict) -> list[str]:
    """본문 비교. 오류 응답은 바이트 그대로, 견적서는 셀 단위로 본다."""
    if want["status"] != 200:
        # 구문 오류의 뒷부분은 파서가 내는 설명이라 구현마다 다르다. 앞의
        # 사용자 문구는 같아야 한다 (파이썬 테스트도 그 앞부분만 본다).
        prefix = "XML을 로드하는중 장애 발생. 장애코드: "
        left = json.loads(want["body"])["error"]
        right = json.loads(got["body"])["error"]
        if (left["code"] == right["code"]
                and left["message"].startswith(prefix)
                and right["message"].startswith(prefix)):
            return []
        if want["body"] != got["body"]:
            return [f"{label}: 오류 본문이 다릅니다\n"
                    f"    파이썬: {want['body'].decode('utf-8', 'replace')}\n"
                    f"    Rust  : {got['body'].decode('utf-8', 'replace')}"]
        return []

    left = openpyxl.load_workbook(io.BytesIO(want["body"]))
    right = openpyxl.load_workbook(io.BytesIO(got["body"]))
    if left.sheetnames != right.sheetnames:
        return [f"{label}: 시트 차례 {left.sheetnames} vs {right.sheetnames}"]
    problems: list[str] = []
    for sheet in left.sheetnames:
        problems.extend(workbook.sheet_problems(label, sheet, left[sheet], right[sheet]))
    return problems[:10]


def compare(label: str, name: str, content_type: str, content: bytes,
            body_path: Path) -> list[str]:
    want = python_response(name, content_type, content)
    got = rust_response(name, content_type, content, body_path)

    problems: list[str] = []
    if want["status"] != got["status"]:
        problems.append(f"{label}: 상태 {want['status']} vs {got['status']}")
    for key in sorted(set(want["headers"]) | set(got["headers"])):
        if key == "Content-Length":
            # 두 라이브러리의 ZIP 직렬화 차이로 파일 크기가 다르다 (결정 0005).
            # 값이 서로 같아야 하는 헤더가 아니라 **제 본문 길이와 맞아야 하는**
            # 헤더다. 양쪽 다 그런지 본다.
            for side, response in (("파이썬", want), ("Rust", got)):
                length = response["headers"].get(key)
                if length is not None and int(length) != len(response["body"]):
                    problems.append(f"{label}: {side} 의 Content-Length 가 본문 길이와 "
                                    f"다릅니다 ({length} vs {len(response['body'])})")
            continue
        if want["headers"].get(key) != got["headers"].get(key):
            problems.append(f"{label}: 헤더 {key} — "
                            f"파이썬 {want['headers'].get(key)!r} / "
                            f"Rust {got['headers'].get(key)!r}")
    for key in sorted(set(want["log"]) | set(got["log"])):
        if want["log"].get(key) != got["log"].get(key):
            problems.append(f"{label}: 로그 {key} — "
                            f"파이썬 {want['log'].get(key)!r} / "
                            f"Rust {got['log'].get(key)!r}")
    problems.extend(body_problems(label, want, got))
    return problems


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    passing = accepted()
    cases = [(f"{name} [{content_type or '없음'}]", name, content_type, content)
             for name, content_type, content in passing]
    cases += [(f"거절: {name}", name, content_type, content)
              for name, content_type, content in REJECTED]

    problems: list[str] = []
    with tempfile.TemporaryDirectory() as folder:
        body_path = Path(folder) / "body.bin"
        for label, name, content_type, content in cases:
            problems.extend(compare(label, name, content_type, content, body_path))

    print(f"대조한 업로드 {len(cases)}건 "
          f"(통과 {len(passing)}, 거절 {len(REJECTED)}) — 상태·헤더·로그·본문")
    if not problems:
        print("파이썬과 Rust 의 응답이 모두 같습니다.")
        return 0
    print(f"다른 곳 {len(problems)}건:")
    for line in problems[:30]:
        print(f"  {line}")
    if len(problems) > 30:
        print(f"  ... 그리고 {len(problems) - 30}건 더")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
