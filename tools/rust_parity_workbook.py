"""파이썬과 Rust 가 만든 **견적서 파일**을 셀 단위로 대조한다.

계획 §Phase 4 의 동등성 판정. 승인 기준은
[결정 0005](../doc/decisions/0005-accept-meaningful-xlsx-parity.md) 다.

1. 제품 정보·수량·단가·금액·수식·시트 순서·병합이 같은가.
2. 첫 페이지(TOTAL)의 로고·머리글 도형이 보존되는가.

여기서 보는 것은 ZIP 바이트가 아니라 **Excel 이 읽는 내용**이다. 두 라이브러리
(openpyxl / umya-spreadsheet)의 직렬화 차이는 진단으로만 남기고 중단 사유로
삼지 않는다.

    python tools/rust_parity_workbook.py            # 두 경로 모두
    python tools/rust_parity_workbook.py --probe    # cargo 탐침만
    python tools/rust_parity_workbook.py --pyo3     # 확장 모듈만

두 경로는 같은 코어를 서로 다른 배포 형태로 부른다. `probe` 는 브라우저 WASM 과
같은 코드이고, `pyo3` 는 데스크톱·CI 가 쓸 확장 모듈이다.
"""
from __future__ import annotations

import datetime as dt
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from quotation.core import resources, xml_reader  # noqa: E402
from quotation.core.writer import ibm_writer  # noqa: E402

PROBE = ("cargo", "run", "--quiet", "--package", "quotation-parity",
         "--bin", "workbook-probe", "--")

FIXTURES = ROOT / "tests" / "fixtures" / "public"

#: 날짜를 고정해야 두 산출물이 같은 값을 갖는다.
TODAY = dt.date(2026, 8, 27)

#: 셀에서 비교하는 것.
CELL_FIELDS = ("value", "number_format", "font", "fill", "alignment", "border")


def documents() -> dict[str, bytes]:
    found = {path.name: path.read_bytes() for path in sorted(FIXTURES.glob("*.xml"))}
    for path in sorted(ROOT.glob("*.xml")):
        found[path.name] = path.read_bytes()
    return found


def python_workbook(data: bytes, template: bytes) -> bytes:
    quotation = xml_reader.parse_bytes(data)
    return ibm_writer.build_bytes(quotation, template, today=TODAY)


def pyo3_workbook(data: bytes, template_path: Path, out: Path) -> bytes:
    """확장 모듈로 만든 견적서 (데스크톱·CI 가 쓸 경로)."""
    import quotation_rust

    del out
    return bytes(quotation_rust.convert_bytes(
        data, template_path.read_bytes(), TODAY.year, TODAY.month, TODAY.day))


def rust_workbook(data: bytes, template_path: Path, out: Path) -> bytes:
    result = subprocess.run(
        [*PROBE, str(template_path), str(out), TODAY.isoformat()],
        cwd=ROOT, input=data, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if result.returncode != 0:
        sys.stderr.write(result.stderr.decode("utf-8", "replace"))
        raise SystemExit(f"Rust 탐침이 {result.returncode} 로 끝났습니다.")
    return out.read_bytes()


# --- 셀 비교 ------------------------------------------------------------------

from openpyxl.styles.colors import COLOR_INDEX  # noqa: E402

#: 시스템 색 — 자동(automatic). 값이 아니라 "지정하지 않음"과 같은 뜻이다.
SYSTEM_INDEXES = {64, 65}


def _color(color) -> str | None:
    """색을 한 가지 표기로 모은다.

    레거시 팔레트 색인과 ARGB 는 **같은 색을 두 가지로 적은 것**이다.
    openpyxl 은 `FFCCFFFF` 로, umya 는 팔레트에 있으면 `indexed="27"` 로
    적는다. 화면에 나오는 색은 같으므로 여기서 한 표기로 모은다.

    색인 64/65(자동)는 색을 지정하지 않은 것과 같으므로 `None` 으로 본다.
    """
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb" and color.rgb:
        return str(color.rgb)
    if kind == "indexed":
        index = color.indexed
        if index in SYSTEM_INDEXES:
            return None
        if 0 <= index < len(COLOR_INDEX):
            return f"FF{COLOR_INDEX[index][2:]}"
        return f"indexed:{index}"
    if kind == "theme":
        return f"theme:{color.theme}"
    return None


def _font(cell) -> tuple:
    font = cell.font
    return (font.name, float(font.sz) if font.sz else None, bool(font.b),
            _color(font.color))


def _fill(cell) -> tuple:
    fill = cell.fill
    pattern = fill.patternType
    if pattern is None:
        return (None, None)
    return (pattern, _color(fill.fgColor))


def _alignment(cell) -> tuple:
    align = cell.alignment
    return (align.horizontal, align.vertical, bool(align.wrap_text))


def _border(cell) -> tuple:
    border = cell.border
    return tuple(
        (getattr(border, side).style, _color(getattr(border, side).color))
        for side in ("left", "right", "top", "bottom")
    )


def _value(cell):
    if isinstance(cell.value, float) and cell.value.is_integer():
        return int(cell.value)
    return cell.value


def cell_facts(cell) -> dict:
    return {
        "value": _value(cell),
        "number_format": cell.number_format,
        "font": _font(cell),
        "fill": _fill(cell),
        "alignment": _alignment(cell),
        "border": _border(cell),
    }


def sheet_problems(name: str, sheet_name: str, left, right) -> list[str]:
    problems: list[str] = []
    if left.max_row != right.max_row or left.max_column != right.max_column:
        problems.append(
            f"{name}[{sheet_name}]: 크기 {left.max_row}x{left.max_column} vs "
            f"{right.max_row}x{right.max_column}")
    merges = (sorted(str(r) for r in left.merged_cells.ranges),
              sorted(str(r) for r in right.merged_cells.ranges))
    if merges[0] != merges[1]:
        only_python = set(merges[0]) - set(merges[1])
        only_rust = set(merges[1]) - set(merges[0])
        problems.append(f"{name}[{sheet_name}]: 병합이 다릅니다 "
                        f"(파이썬만 {sorted(only_python)}, Rust만 {sorted(only_rust)})")
    if left.print_area != right.print_area:
        problems.append(f"{name}[{sheet_name}]: 인쇄 영역 "
                        f"{left.print_area} vs {right.print_area}")

    rows = max(left.max_row, right.max_row)
    columns = max(left.max_column, right.max_column)
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            want = cell_facts(left.cell(row=row, column=column))
            got = cell_facts(right.cell(row=row, column=column))
            for field in CELL_FIELDS:
                if want[field] != got[field]:
                    coordinate = left.cell(row=row, column=column).coordinate
                    problems.append(
                        f"{name}[{sheet_name}]{coordinate}.{field}: "
                        f"파이썬 {want[field]!r} / Rust {got[field]!r}")
    for row in range(1, rows + 1):
        want = left.row_dimensions[row].height
        got = right.row_dimensions[row].height
        if want != got:
            problems.append(f"{name}[{sheet_name}] {row}행 높이: {want} vs {got}")
    return problems


def drawing_problems(name: str, template: bytes, python_bytes: bytes,
                     rust_bytes: bytes) -> list[str]:
    """첫 페이지 도형이 템플릿 그대로 남아 있는가 (결정 0005 두 번째 조건).

    보는 것은 셋이다.

    1. 두 산출물이 **같은 그림 부품**을 담고 있는가 (openpyxl 은 저장할 때
       그림을 버리고 TOTAL 시트 것만 다시 붙인다).
    2. 그 부품의 바이트가 서로 같은가.
    3. 그 부품이 템플릿 원본과 바이트까지 같은가.
    """
    def drawing_parts(data: bytes) -> dict[str, bytes]:
        with zipfile.ZipFile(io_bytes(data)) as z:
            return {n: z.read(n) for n in z.namelist()
                    if n.startswith(("xl/drawings/", "xl/media/"))}

    problems: list[str] = []
    want = drawing_parts(python_bytes)
    got = drawing_parts(rust_bytes)
    source = drawing_parts(template)

    if set(want) != set(got):
        only_python = sorted(set(want) - set(got))
        only_rust = sorted(set(got) - set(want))
        problems.append(f"{name}: 그림 부품이 다릅니다 "
                        f"(파이썬만 {only_python}, Rust만 {only_rust})")
    for part in sorted(set(want) & set(got)):
        if want[part] != got[part]:
            problems.append(f"{name}: {part} 가 파이썬 산출물과 다릅니다")
        elif part in source and source[part] != got[part]:
            problems.append(f"{name}: {part} 가 템플릿과 다릅니다")
    return problems


def io_bytes(data: bytes):
    import io
    return io.BytesIO(data)


#: 같은 코어를 부르는 두 배포 형태.
BACKENDS = {
    "probe": ("cargo 탐침", rust_workbook),
    "pyo3": ("확장 모듈 quotation_rust", pyo3_workbook),
}


def compare(name: str, data: bytes, template_bytes: bytes, template_path: Path,
            out: Path, produce) -> list[str]:
    python_bytes = python_workbook(data, template_bytes)
    rust_bytes = produce(data, template_path, out)

    left = openpyxl.load_workbook(io_bytes(python_bytes))
    right = openpyxl.load_workbook(io_bytes(rust_bytes))

    problems: list[str] = []
    if left.sheetnames != right.sheetnames:
        return [f"{name}: 시트 차례 {left.sheetnames} vs {right.sheetnames}"]
    for sheet_name in left.sheetnames:
        if left[sheet_name].sheet_state != right[sheet_name].sheet_state:
            problems.append(f"{name}[{sheet_name}]: 표시 상태 "
                            f"{left[sheet_name].sheet_state} vs "
                            f"{right[sheet_name].sheet_state}")
        problems.extend(sheet_problems(name, sheet_name,
                                       left[sheet_name], right[sheet_name]))
    problems.extend(drawing_problems(name, template_bytes, python_bytes, rust_bytes))
    return problems


def main(argv: list[str]) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    template_path = resources.default_template_path()
    template_bytes = template_path.read_bytes()
    chosen = [name for name in BACKENDS if f"--{name}" in argv] or list(BACKENDS)

    found = documents()
    print(f"대조한 문서 {len(found)}건 — 셀 값·수식·서식·병합·인쇄영역·도형")
    failed = 0
    with tempfile.TemporaryDirectory() as folder:
        out = Path(folder) / "rust.xlsx"
        for backend in chosen:
            label, produce = BACKENDS[backend]
            problems: list[str] = []
            try:
                for name, data in found.items():
                    problems.extend(
                        compare(name, data, template_bytes, template_path, out, produce))
            except ImportError as error:
                print(f"[{backend}] {label}: 부를 수 없습니다 ({error})")
                failed = 1
                continue
            if not problems:
                print(f"[{backend}] {label}: 파이썬과 견적서가 모두 같습니다.")
                continue
            failed = 1
            print(f"[{backend}] {label}: 다른 곳 {len(problems)}건")
            for line in problems[:40]:
                print(f"  {line}")
            if len(problems) > 40:
                print(f"  ... 그리고 {len(problems) - 40}건 더")
    return failed


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
