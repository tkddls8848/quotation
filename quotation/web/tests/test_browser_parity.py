"""브라우저 변환이 데스크톱 변환과 같은 결과를 내는지 대조한다.

무료 계정에서는 변환이 브라우저에서 돈다. **결과가 조금이라도 달라지면 안 된다**
— 금액 한 자리, 품목 한 줄, 서식 한 칸도. 그래서 화면이 쓰는 엔진
(`quotation/web/src/engine.js`)을 Node 에서 **같은 파일 그대로** 돌려 나온
바이트를, 데스크톱이 쓰는 경로(확장 모듈 `quotation_rust`)의 산출물과 대조한다.

두 경로는 같은 Rust 코어를 서로 다른 배포 형태로 부른다 — 한쪽은 wasm, 다른
한쪽은 네이티브 확장이다. 여기서 잡히는 것은 규칙이 아니라 **포장과 경계**의
문제다 (자산이 낡았거나, 바이트가 잘렸거나, 날짜·모드가 다르게 넘어갔거나).
규칙 자체는 `rust/core` 의 테스트와 `tests/` 의 골든 회귀가 지킨다.

대조 방식은 셋이다.

    1. openpyxl 로 열어 셀 단위로 비교한다 (`quotation/tools/xlsx_content.py`).
    2. .xlsx(zip) 의 부품 목록이 같은지 본다.
    3. 첫 페이지의 로고·도형이 템플릿 원본과 바이트까지 같은지 본다.

거기에 화면과의 계약(상태 코드·헤더·로그·거절 문구)을 함께 본다.

엔진 자산이 없거나 node 가 없으면 건너뛴다.

    python quotation/web/scripts/build_browser_engine.py
"""
from __future__ import annotations

import datetime as dt
import json
import shutil
import subprocess
import zipfile
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

import xlsx_content
from quotation.core import convert

FEATURE = Path(__file__).resolve().parents[2]   # quotation/
ROOT = FEATURE.parent
ENGINE_DIR = ROOT / "web" / "public" / "engine"
DRIVER = FEATURE / "web" / "scripts" / "browser_convert.mjs"

#: 대조할 파일. 어느 쪽 문서(IBM/레노버)든 사람이 고르지 않고 내용으로
#: 갈린다 — 브라우저와 데스크톱이 같은 판단을 내려야 한다.
CASES = (
    "new_quote.xml",
    "upgrade_quote.xml",
    "no_charge.xml",
    "euckr_quote.xml",
    # 통합(레노버) 문서도 브라우저와 데스크톱이 같은 바이트를 내야 한다.
    "integrated_quote.xml",
    # 요약표를 푸는 길(base64+gzip)도 브라우저에서 같아야 한다.
    "integrated_summary_quote.xml",
)

#: 실패도 같아야 한다. (이름, 내용, 기대 상태, 기대 오류 코드)
ERROR_CASES = (
    ("broken.xml", b"<CFXML><CFData><ProductLineItem>", 422, "INVALID_QUOTATION_XML"),
    ("empty.xml", b"   \n", 400, "INVALID_REQUEST"),
    ("notes.txt", "XML 이 아닙니다".encode("utf-8"), 415, "UNSUPPORTED_MEDIA_TYPE"),
)

pytestmark = pytest.mark.skipif(
    not (ENGINE_DIR / "engine.json").is_file() or shutil.which("node") is None,
    reason="브라우저 엔진 자산이 없습니다. "
           "python quotation/web/scripts/build_browser_engine.py 를 먼저 실행하십시오",
)


@pytest.fixture(scope="module")
def browser_results(tmp_path_factory, fixtures) -> dict[str, dict]:
    """엔진을 한 번 띄워 모든 사례를 변환한다."""
    out = tmp_path_factory.mktemp("browser")
    bad = tmp_path_factory.mktemp("bad")
    for name, content, _, _ in ERROR_CASES:
        (bad / name).write_bytes(content)

    inputs = [str(fixtures / name) for name in CASES]
    inputs += [str(bad / name) for name, _, _, _ in ERROR_CASES]

    proc = subprocess.run(
        ["node", str(DRIVER), str(ENGINE_DIR), str(out), *inputs],
        capture_output=True, text=True, timeout=600)
    assert proc.returncode == 0, f"브라우저 엔진 실행 실패:\n{proc.stderr[-4000:]}"

    results = {}
    for name in [*CASES, *(n for n, _, _, _ in ERROR_CASES)]:
        stem = Path(name).stem
        meta = json.loads((out / f"{stem}.json").read_text(encoding="utf-8"))
        xlsx = out / f"{stem}.xlsx"
        meta["xlsx"] = xlsx.read_bytes() if xlsx.is_file() else b""
        results[name] = meta
    return results


def _desktop(fixtures, name: str, today: dt.date) -> bytes:
    """데스크톱이 쓰는 경로. 같은 Rust 코어를 확장 모듈로 부른다."""
    return convert.convert_bytes(
        (fixtures / name).read_bytes(), today=today, source_name=name).xlsx


# --- 내용 대조 -----------------------------------------------------------------

@pytest.mark.parametrize("name", CASES)
def test_browser_xlsx_matches_the_desktop_cell_by_cell(name, browser_results,
                                                       fixtures):
    """값·수식·서식·글꼴·채우기·테두리·정렬·병합·인쇄영역·행높이."""
    result = browser_results[name]
    assert result["status"] == 200, result.get("body_utf8")

    expected = _desktop(fixtures, name, dt.date.fromisoformat(result["today"]))
    want = openpyxl.load_workbook(BytesIO(expected))
    got = openpyxl.load_workbook(BytesIO(result["xlsx"]))
    assert want.sheetnames == got.sheetnames

    problems = []
    for sheet in want.sheetnames:
        problems.extend(
            xlsx_content.sheet_problems(name, sheet, want[sheet], got[sheet]))
    assert not problems, "\n".join(problems[:20])


@pytest.mark.parametrize("name", CASES)
def test_browser_xlsx_ships_the_same_parts(name, browser_results, fixtures):
    """부품 목록이 같아야 한다. 빠지는 것도, 남는 것도 없어야 한다."""
    result = browser_results[name]
    expected = _desktop(fixtures, name, dt.date.fromisoformat(result["today"]))

    desktop = zipfile.ZipFile(BytesIO(expected))
    browser = zipfile.ZipFile(BytesIO(result["xlsx"]))
    with desktop, browser:
        assert sorted(desktop.namelist()) == sorted(browser.namelist())


def test_browser_keeps_the_template_drawings(browser_results, template_bytes):
    """로고와 머리말 도형이 템플릿 원본 그대로 실려 나오는지 (양식 뒤틀림 방지)."""
    produced = zipfile.ZipFile(BytesIO(browser_results["new_quote.xml"]["xlsx"]))
    template = zipfile.ZipFile(BytesIO(template_bytes))
    with produced, template:
        parts = [name for name in produced.namelist()
                 if name.startswith(("xl/media/", "xl/drawings/"))]
        assert parts
        for part in parts:
            assert produced.read(part) == template.read(part), part


# --- 화면과의 계약 --------------------------------------------------------------

@pytest.mark.parametrize("name", CASES)
def test_browser_returns_the_contract_headers_and_log(name, browser_results):
    """상태 코드·다운로드 이름·템플릿 판본·로그 항목."""
    result = browser_results[name]
    stem = Path(name).stem

    assert result["status"] == 200
    assert result["headers"]["Content-Type"].endswith("spreadsheetml.sheet")
    assert f'filename="{stem}.xlsx"' in result["headers"]["Content-Disposition"]
    assert result["headers"]["X-Template-Version"].startswith("sha256-")
    assert result["headers"]["Content-Length"] == str(len(result["xlsx"]))

    log = result["log"]
    assert log["outcome"] == "ok"
    assert log["status"] == 200
    assert log["mode"] in ("unix", "integrated")
    assert log["group_count"] >= 1
    assert log["line_count"] >= log["group_count"]
    assert log["template_version"] == result["headers"]["X-Template-Version"]
    for field in ("input_size_bucket", "output_size_bucket"):
        assert log[field].startswith("<=")


@pytest.mark.parametrize("name,content,status,code", ERROR_CASES)
def test_browser_rejects_the_same_way_the_contract_says(name, content, status,
                                                        code, browser_results):
    """거절도 계약대로여야 한다. 같은 상태 코드, 같은 오류 코드, 사람이 읽을 문구."""
    result = browser_results[name]
    assert result["status"] == status

    error = json.loads(result["body_utf8"])["error"]
    assert error["code"] == code
    assert error["message"]
    assert error["request_id"]
    assert result["log"]["outcome"] == "error"
    assert result["log"]["error_code"] == code


def test_browser_stamps_the_seoul_date(browser_results):
    """견적 날짜는 브라우저의 지역 시간이 아니라 Asia/Seoul 기준이다."""
    seoul = dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).date()
    stamped = browser_results["new_quote.xml"]["today"]
    assert dt.date.fromisoformat(stamped) == seoul
