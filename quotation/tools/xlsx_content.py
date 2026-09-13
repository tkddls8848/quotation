"""견적서 두 벌이 **Excel 이 읽는 내용**으로 같은지 본다.

무엇을 같게 볼지는 [결정 0005](../doc/decisions/0005-accept-meaningful-xlsx-parity.md)
와 [결정 0009](../doc/decisions/0009-browser-parity-judges-content.md)가 정한다.
ZIP 바이트가 아니라 값·수식·서식·글꼴·채우기·테두리·정렬·병합·인쇄영역·행높이다.
색은 팔레트 색인과 ARGB 를 같은 색으로 본다 — OOXML 을 적는 라이브러리마다
고르는 표기가 다르기 때문이다.

    from xlsx_content import sheet_problems
"""
from __future__ import annotations

from openpyxl.styles.colors import COLOR_INDEX

#: 셀에서 비교하는 것.
CELL_FIELDS = ("value", "number_format", "font", "fill", "alignment", "border")  # noqa: E402

#: 시스템 색 — 자동(automatic). 값이 아니라 "지정하지 않음"과 같은 뜻이다.
SYSTEM_INDEXES = {64, 65}


def _color(color) -> str | None:
    """색을 한 가지 표기로 모은다.

    레거시 팔레트 색인과 ARGB 는 **같은 색을 두 가지로 적은 것**이다.
    openpyxl 은 `FFCCFFFF` 로, umya 는 팔레트에 있으면 `indexed="27"` 로
    적는다. 화면에 나오는 색은 같으므로 여기서 한 표기로 모은다.

    색인 64/65(자동)는 색을 지정하지 않은 것과 같으므로 `None` 으로 본다.
    """
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb" and color.rgb:
        return str(color.rgb)
    if kind == "indexed":
        index = color.indexed
        if index in SYSTEM_INDEXES:
            return None
        if 0 <= index < len(COLOR_INDEX):
            return f"FF{COLOR_INDEX[index][2:]}"
        return f"indexed:{index}"
    if kind == "theme":
        return f"theme:{color.theme}"
    return None


def same_color(color, argb: str) -> bool:
    """같은 색인가. 팔레트 색인과 ARGB 는 같은 색을 두 가지로 적은 것이다."""
    return _color(color) == argb


def _font(cell) -> tuple:
    font = cell.font
    return (font.name, float(font.sz) if font.sz else None, bool(font.b),
            _color(font.color))


def _fill(cell) -> tuple:
    fill = cell.fill
    pattern = fill.patternType
    if pattern is None:
        return (None, None)
    return (pattern, _color(fill.fgColor))


def _alignment(cell) -> tuple:
    align = cell.alignment
    return (align.horizontal, align.vertical, bool(align.wrap_text))


def _border(cell) -> tuple:
    border = cell.border
    return tuple(
        (getattr(border, side).style, _color(getattr(border, side).color))
        for side in ("left", "right", "top", "bottom")
    )


def _value(cell):
    if isinstance(cell.value, float) and cell.value.is_integer():
        return int(cell.value)
    return cell.value


def cell_facts(cell) -> dict:
    return {
        "value": _value(cell),
        "number_format": cell.number_format,
        "font": _font(cell),
        "fill": _fill(cell),
        "alignment": _alignment(cell),
        "border": _border(cell),
    }


def sheet_problems(name: str, sheet_name: str, left, right) -> list[str]:
    problems: list[str] = []
    if left.max_row != right.max_row or left.max_column != right.max_column:
        problems.append(
            f"{name}[{sheet_name}]: 크기 {left.max_row}x{left.max_column} vs "
            f"{right.max_row}x{right.max_column}")
    merges = (sorted(str(r) for r in left.merged_cells.ranges),
              sorted(str(r) for r in right.merged_cells.ranges))
    if merges[0] != merges[1]:
        only_python = set(merges[0]) - set(merges[1])
        only_rust = set(merges[1]) - set(merges[0])
        problems.append(f"{name}[{sheet_name}]: 병합이 다릅니다 "
                        f"(파이썬만 {sorted(only_python)}, Rust만 {sorted(only_rust)})")
    if left.print_area != right.print_area:
        problems.append(f"{name}[{sheet_name}]: 인쇄 영역 "
                        f"{left.print_area} vs {right.print_area}")

    rows = max(left.max_row, right.max_row)
    columns = max(left.max_column, right.max_column)
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            want = cell_facts(left.cell(row=row, column=column))
            got = cell_facts(right.cell(row=row, column=column))
            for field in CELL_FIELDS:
                if want[field] != got[field]:
                    coordinate = left.cell(row=row, column=column).coordinate
                    problems.append(
                        f"{name}[{sheet_name}]{coordinate}.{field}: "
                        f"파이썬 {want[field]!r} / Rust {got[field]!r}")
    for row in range(1, rows + 1):
        want = left.row_dimensions[row].height
        got = right.row_dimensions[row].height
        if want != got:
            problems.append(f"{name}[{sheet_name}] {row}행 높이: {want} vs {got}")
    return problems
